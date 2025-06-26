from flask import Flask, render_template, request, send_file, redirect, url_for, flash
from docx import Document
import os
from datetime import datetime

app = Flask(__name__)
app.secret_key = 'supersecret'
OUTPUT_FOLDER = 'generated_docs'
TEMPLATE_PATH = os.path.join('doc_templates', 'protocol_template.docx')
os.makedirs(OUTPUT_FOLDER, exist_ok=True)

@app.route('/', methods=['GET', 'POST'])
def index():
    if request.method == 'POST':
        system = request.form['system']
        author = request.form['author']
        date = request.form['date']
        description = request.form['description']
        test_steps = request.form['test_steps']

        # Load template
        doc = Document(TEMPLATE_PATH)

        # Replace placeholders
        for p in doc.paragraphs:
            if '[SYSTEM]' in p.text:
                p.text = p.text.replace('[SYSTEM]', system)
            if '[AUTHOR]' in p.text:
                p.text = p.text.replace('[AUTHOR]', author)
            if '[DATE]' in p.text:
                p.text = p.text.replace('[DATE]', date)
            if '[DESCRIPTION]' in p.text:
                p.text = p.text.replace('[DESCRIPTION]', description)
            if '[TEST_STEPS]' in p.text:
                p.text = p.text.replace('[TEST_STEPS]', test_steps)

        # Save generated doc
        out_filename = f"{system}_protocol_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
        out_path = os.path.join(OUTPUT_FOLDER, out_filename)
        doc.save(out_path)

        return render_template('success.html', filename=out_filename)
    return render_template('index.html')

@app.route('/download/<filename>')
def download(filename):
    path = os.path.join(OUTPUT_FOLDER, filename)
    if os.path.exists(path):
        return send_file(path, as_attachment=True)
    flash("File not found.", "danger")
    return redirect(url_for('index'))

if __name__ == '__main__':
    app.run(debug=True)
